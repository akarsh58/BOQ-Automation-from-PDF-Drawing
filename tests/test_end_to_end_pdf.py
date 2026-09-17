from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app


def test_real_pdf_generates_preliminary_workbook():
    drawing = Path(__file__).parents[1] / "sample_input" / "sample_floor_plan.pdf"
    assert drawing.exists()
    with TestClient(app) as client, drawing.open("rb") as handle:
        response = client.post(
            "/generate",
            files={"file": (drawing.name, handle, "application/pdf")},
        )
    assert response.status_code == 200, response.text
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    try:
        assert "BOQ" in workbook.sheetnames
        assert "Takeoff" in workbook.sheetnames
        assert "QA" in workbook.sheetnames
        assert workbook["Cover"]["A1"].value == "PRELIMINARY ESTIMATED BOQ"
        assert workbook["QA"]["B2"].value == "AUTOMATED_PRELIMINARY"
    finally:
        workbook.close()
