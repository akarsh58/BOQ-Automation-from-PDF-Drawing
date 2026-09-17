import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from app.boq_engine import load_knowledge_base
from app.main import _automatic_takeoff
from app.tender_export import write_tender_workbook


class AutomaticPipelineTests(unittest.TestCase):
    def test_rooms_keep_page_positions_and_confidence(self):
        pages = [
            {
                "page": 0,
                "px_per_unit": 100,
                "scale_confidence": 0.95,
                "scale_method": "text_scale",
                "rooms": [{"id": "r1", "x": 200, "y": 300, "width": 400, "height": 500, "name": "Kitchen", "confidence": 0.8}],
            }
        ]
        doc = _automatic_takeoff(pages, "plan.pdf", 3, 0.23, [100], 0.95, "text_scale", "")
        room = next(element for element in doc.elements if element.type == "space")
        self.assertEqual(room.polygon_m[0], (2.0, 3.0))
        self.assertEqual(room.extra["detection_confidence"], 0.8)
        self.assertTrue(doc.automatic)
        self.assertFalse(doc.qs_signed)
        self.assertFalse(doc.ifc_units_confirmed)

    def test_automatic_workbook_is_preliminary_and_flags_scale(self):
        pages = [
            {
                "page": 0,
                "px_per_unit": 50,
                "scale_confidence": 0.0,
                "scale_method": "missing",
                "rooms": [{"id": "r1", "x": 20, "y": 30, "width": 200, "height": 200, "name": "Room 1", "confidence": 0.0}],
            }
        ]
        doc = _automatic_takeoff(
            pages, "plan.pdf", 3, 0.23, [50], 0.0, "missing",
            "No reliable scale text; fallback is preliminary.",
        )
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "preliminary.xlsx"
            write_tender_workbook(output, doc, load_knowledge_base(), enforce_signoff=False)
            workbook = load_workbook(BytesIO(output.read_bytes()))
            cover_values = [row[1] for row in workbook["Cover"].iter_rows(min_col=1, max_col=2, values_only=True)]
            deficiencies = list(workbook["Deficiencies"].values)
            cover_title = workbook["Cover"]["A1"].value
            workbook.close()
            self.assertIn("PRELIMINARY ESTIMATED BOQ", cover_title)
            self.assertIn("Automatic preliminary - QS review required", cover_values)
            self.assertIn("Not supplied - preliminary only", cover_values)
            self.assertTrue(any("low-confidence" in str(row[1]) for row in deficiencies[1:]))


if __name__ == "__main__":
    unittest.main()
